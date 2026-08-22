from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION
from pathlib import Path

OUT = Path('/Users/elokuu/Desktop/一生足迹/一生足迹_产品设计说明书_v1.0.docx')

INK = '17202A'
NAVY = '18324A'
ACCENT = 'EF3340'
MUTED = '667085'
LIGHT = 'EEF2F6'
PALE_RED = 'FDECEF'
GREEN = '16794A'
AMBER = '9A6700'
WHITE = 'FFFFFF'

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.78)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'STHeiti'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'STHeiti')
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.22

for name, size, color, before, after in [
    ('Title', 26, NAVY, 0, 6),
    ('Subtitle', 12, MUTED, 0, 16),
    ('Heading 1', 17, NAVY, 18, 8),
    ('Heading 2', 13, NAVY, 13, 6),
    ('Heading 3', 11.5, NAVY, 9, 4),
]:
    st = styles[name]
    st.font.name = 'STHeiti'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'STHeiti')
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = name != 'Subtitle'
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for list_name in ['List Bullet', 'List Number']:
    st = styles[list_name]
    st.font.name = 'STHeiti'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'STHeiti')
    st.font.size = Pt(10.5)
    st.paragraph_format.left_indent = Inches(0.38)
    st.paragraph_format.first_line_indent = Inches(-0.19)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.2

if 'Status Label' not in styles:
    s = styles.add_style('Status Label', WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = 'STHeiti'; s.font.size = Pt(9); s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(MUTED)
    s.paragraph_format.space_before = Pt(3); s.paragraph_format.space_after = Pt(3)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None: node = OxmlElement('w:' + m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None: tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'), 'dxa')
    tblLayout = OxmlElement('w:tblLayout'); tblLayout.set(qn('w:type'), 'fixed'); tblPr.append(tblLayout)
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        col = OxmlElement('w:gridCol'); col.set(qn('w:w'), str(w)); grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tcW = cell._tc.get_or_add_tcPr().find(qn('w:tcW'))
            tcW.set(qn('w:w'), str(widths[idx])); tcW.set(qn('w:type'), 'dxa')
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); repeat = OxmlElement('w:tblHeader'); repeat.set(qn('w:val'), 'true'); trPr.append(repeat)

def add_table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'; set_table_widths(t, widths)
    for i, h in enumerate(headers):
        shade(t.rows[0].cells[i], NAVY)
        p = t.rows[0].cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h); r.bold = True; r.font.color.rgb = RGBColor.from_string(WHITE); r.font.size = Pt(9.5)
    set_repeat_header(t.rows[0])
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            if ridx % 2: shade(cells[i], 'F7F9FB')
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(value)); r.font.size = Pt(9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def add_bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.add_run(text); return p

def add_number(text):
    p = doc.add_paragraph(style='List Number'); p.add_run(text); return p

def add_status(text, status):
    colors = {'已实现': GREEN, '部分实现': AMBER, '规划中': MUTED, '需真机验证': ACCENT}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(5)
    a = p.add_run(status + '  '); a.bold = True; a.font.size = Pt(9); a.font.color.rgb = RGBColor.from_string(colors[status])
    b = p.add_run(text); b.font.size = Pt(10.3)

def add_callout(title, body, fill=LIGHT):
    t = doc.add_table(rows=1, cols=1); set_table_widths(t, [9360]); shade(t.cell(0,0), fill)
    p = t.cell(0,0).paragraphs[0]; p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title); r.bold = True; r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = t.cell(0,0).add_paragraph(body); p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

# Running furniture
hp = sec.header.paragraphs[0]
hp.text = '一生足迹  ·  产品设计说明书'
hp.style = styles['Status Label']
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = fp.add_run('内部产品文档  |  v1.0  |  2026-08-17'); fr.font.size = Pt(8.5); fr.font.color.rgb = RGBColor.from_string(MUTED)

# Cover / masthead
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(7)
r = p.add_run('PRODUCT SPECIFICATION'); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(ACCENT)
doc.add_paragraph('一生足迹', style='Title')
doc.add_paragraph('地图足迹记录与照片回忆产品设计说明书', style='Subtitle')
add_table(['文档字段', '内容'], [
    ('版本', 'v1.0（当前设计与实现汇总）'),
    ('平台', 'iOS 17+，SwiftUI / MapKit / Photos / SwiftData'),
    ('目标设备', 'iPhone，当前真机测试设备为 iPhone 15 Pro Max'),
    ('文档日期', '2026 年 8 月 17 日'),
    ('状态口径', '已实现 / 部分实现 / 规划中 / 需真机验证'),
], [2100, 7260])
add_callout('产品一句话', '以地图为主界面，低功耗持续记录个人移动足迹，并把系统照片按地点、路线和时间组织成可探索的地理回忆。', PALE_RED)

doc.add_heading('1. 产品目标与设计原则', level=1)
add_bullet('地图优先：足迹、路线、照片与当前位置共享同一个空间语境。')
add_bullet('专业运动感：路线清晰、点位明确、深色地图与高对比强调色形成 Garmin / COROS 式信息气质。')
add_bullet('隐私与本地优先：足迹与照片索引主要保存在本机；涉及系统照片删除时必须经过 iOS 系统确认。')
add_bullet('渐进披露：缩放较小时聚合与降噪，放大后才显示细节；浏览进度弱化，不直接强调照片总量。')
add_bullet('性能优先：地图拖动、缩放和照片切换期间减少重复计算、全量解码与视图漂移。')

doc.add_heading('2. 信息架构与导航', level=1)
add_table(['一级入口', '核心内容', '当前状态'], [
    ('地图', '足迹点、路线、照片聚合、定位、地图图层与 2D/3D', '已实现'),
    ('统计', '里程、足迹点、活跃天数等汇总', '已实现；地图页汇总面板已按要求移除'),
    ('设置', '后台记录、地图源、自定义瓦片、语言及其他偏好', '已实现/持续优化'),
    ('照片探索', '由地图照片聚合或行政区域进入沉浸浏览', '已实现'),
], [1500, 5900, 1960])

doc.add_heading('3. 地图体验', level=1)
doc.add_heading('3.1 视觉与图层', level=2)
add_status('深色专业运动风格；路线采用高对比红色主线、描边与方向箭头，足迹点更清晰。', '已实现')
add_status('标准地图、卫星地图与自定义 XYZ/TMS 瓦片地图切换。', '已实现')
add_status('专业等高线地图通过 MapTiler Outdoor v4 等自定义瓦片导入，不保留无效的内置“专业等高线”入口。', '已实现')
add_status('支持 2D 与俯视 3D 视角调整。', '已实现')
add_status('真正面向户外专业用途的矢量等高线分层、坡度阴影、离线瓦片包。', '规划中')

doc.add_heading('3.2 点、线与照片聚合', level=2)
add_bullet('照片标记与路线使用地理坐标绑定，地图拖动时随地图投影移动，而不是作为屏幕浮层漂移。')
add_bullet('路线附近照片优先投影/吸附到路线线段，可按组显示；远离路线的照片保留真实质心。')
add_bullet('不同缩放层级使用聚合与迟滞阈值，避免点位在边界缩放时频繁闪烁。')
add_bullet('地图点击与照片标记点击分离命中区域，降低滑动地图时误触照片。')
add_bullet('隐藏/显示路线和照片图层保留稳定状态，已修复路线恢复闪退及照片开关短时失效问题。')

doc.add_heading('3.3 定位控制', level=2)
add_bullet('单击定位按钮：回到当前定位地点。')
add_bullet('双击定位按钮：进入跟随模式，地图根据陀螺仪/指南针朝向旋转。')
add_bullet('位置、经纬度、海拔、精度与时间信息靠左上对齐，控制区保持紧凑。')

doc.add_heading('4. 足迹记录策略', level=1)
add_status('后台足迹开关启用后使用系统定位能力自动记录。', '已实现')
add_status('本地移动采用重大位移与质量门槛控制，降低低精度漂移和电量消耗。', '已实现')
add_status('高铁环境采用稀疏关键点策略，避免高速移动产生过密点列。', '已实现')
add_status('飞机环境过滤短距离漂移，只保留足以表达长距离移动的关键点。', '已实现')
add_status('后台记录频率不是固定秒数，而是由移动距离、速度、精度和系统唤醒共同决定。', '已实现')
add_status('长时间真实道路、高铁与飞行场景的电量基准测试及自动运动模式识别。', '需真机验证')

doc.add_heading('5. 地图源管理', level=1)
add_bullet('设置中提供“自定义地图源”文件夹，集中管理已授权的自定义瓦片。')
add_bullet('支持选择、切换和删除自定义地图源；冷启动恢复上次选择。')
add_bullet('输入框支持自动识别 iframe、MapTiler 页面地址、XYZ 与 TMS 模板，不保留复制/粘贴图标。')
add_bullet('已验证 MapTiler Outdoor v4 URL 自动补全与导入。')
add_bullet('拒绝非 HTTPS 地址与未经授权的 Google 非官方瓦片直链，避免授权、稳定性与合规风险。')

doc.add_heading('6. 照片地图入口与抽样', level=1)
add_bullet('点击地图照片聚合后跳过中间预览，直接进入照片探索页。')
add_bullet('进入地点或区域时，从完整照片集合中按时间分层随机抽取最多 20 张。')
add_bullet('重点打散相邻日期，避免同一行程照片连续出现；若总数少于 20 张，仅打乱已有照片。')
add_bullet('右上角不显示照片总数量；底部进度条弱化，并在停止交互后自动隐藏。')
add_bullet('一组浏览完成后显示“回顾完毕”；“再来一组”留在当前地点重新抽样，不切换地图区域。')
add_bullet('重新抽样时显示约 0.72 秒“正在洗牌”过渡，随后从第一张开始。')

doc.add_heading('7. 单张沉浸浏览', level=1)
add_table(['能力', '产品规则', '状态'], [
    ('照片布局', '保持原始横竖比例，居中显示并限制在当前窗口宽度内', '已实现'),
    ('环境背景', '使用当前照片颜色模糊延展，横屏照片不出现生硬空白', '已实现'),
    ('左右切换', '照片跟随手指移动并提前露出相邻照片；左滑下一张、右滑上一张', '已实现'),
    ('横竖切换', '切换时使用动画，避免容器直接跳变', '部分实现；仍可加强比例预判'),
    ('底部信息', '具体地点、拍摄时间与海拔；不显示层级名称', '部分实现；“多久以前”待替换'),
    ('双指缩放', '双指手势进入当前照片所在日期的当日照片', '已实现'),
    ('上滑', '明确垂直手势暂存待删除，拖动出现红色蒙层与垃圾桶反馈', '已实现'),
], [1600, 5900, 1860])

doc.add_heading('8. Live Photo', level=1)
add_bullet('Live Photo 左上角显示“实况”标记。')
add_bullet('长按照片手动播放，松开停止。')
add_bullet('单击“实况”标记启用当前会话自动静音播放，再次单击关闭。')
add_bullet('播放完成后恢复左右滑动，避免播放器状态吞掉翻页手势。')
add_bullet('提前准备当前 Live Photo 资源，减少自动播放卡顿；设置页不再提供全局“自动播放 Live Photo”开关。')

doc.add_heading('9. 删除与复核流程', level=1)
add_number('用户在单张浏览中上滑，照片加入“待删除”集合；再次操作可撤销。')
add_number('刷完一组后进入深色复核卡，顶部显示“回顾完毕 / 请确认需要删除的照片”。')
add_number('待删除照片根据数量自动缩放：1 张单列，2–4 张两列，5–9 张三列，10 张以上四列。')
add_number('点击“放弃，再来一组”：保留照片，并在当前地点重新随机抽样。')
add_number('点击“确认删除”：请求 Photos 读写权限并调用 iOS 系统照片删除确认。')
add_number('系统确认后从系统照片库删除；启用 iCloud 照片时同步到其他设备，并进入“最近删除”。')
add_number('用户拒绝或删除失败时，不清理 App 记录，显示错误信息，避免数据状态不一致。')
add_callout('不可绕过的系统规则', '系统照片删除确认弹窗的视觉、文案与按钮由 iOS 管理，App 不能完全自定义。产品只负责弹窗前的复核 UI 与成功/失败后的数据一致性。', LIGHT)

doc.add_heading('10. 结束页与过渡', level=1)
add_bullet('结束页采用高圆角深灰长卡：顶部标题与副标题、中部低对比徽章、底部主操作。')
add_bullet('无待删除照片时主按钮为“再来一组”。')
add_bullet('有待删除照片时底部并排“放弃，再来一组 / 确认删除”。')
add_bullet('所有结束页内容限制在安全区和窗口宽度内，底层页面被深色遮罩压低。')

doc.add_heading('11. 设置与国际化', level=1)
add_bullet('后台足迹记录开关。')
add_bullet('地图源与自定义瓦片管理。')
add_bullet('语言设置：简体中文、繁体中文、英语、法语。')
add_bullet('已移除无效的内置专业等高线地图源。')
add_bullet('当前无需“右滑删除免确认”开关；删除统一采用上滑暂存 + 组末确认 + iOS 系统确认。')

doc.add_heading('12. 性能与稳定性设计', level=1)
add_bullet('首屏使用品牌 Logo 过渡遮罩降低等待感，并将重数据准备延后。')
add_bullet('地图缩放/拖动期间减少标注重复生成与照片解码，修复点线照片跟手漂移和放大缩小卡顿。')
add_bullet('照片页预取当前及相邻图片，区分缩略图与高分辨率加载。')
add_bullet('Live Photo 播放结束主动释放播放状态，避免页面卡死。')
add_bullet('自定义地图源冷启动恢复、瓦片加载失败与慢加载仍需持续真机监测。')

doc.add_heading('13. 当前完成度', level=1)
add_table(['模块', '完成度', '备注'], [
    ('地图与足迹', '高', '核心点线、定位、图层与 2D/3D 已落地'),
    ('后台记录', '中高', '策略已实现，长期功耗与交通模式需积累真机样本'),
    ('地图源', '中高', '自定义导入、管理与 MapTiler 已覆盖'),
    ('照片单张浏览', '高', '横竖布局、手势、Live Photo 与结束页已完成'),
    ('系统照片删除', '中高', '权限与系统确认已接入，需用非重要照片真机验收'),
    ('洗牌拼贴模式', '低', '尚未实现不规则多照片拼贴'),
    ('多语言', '中', '入口与语言范围已确定，需持续检查全量文案覆盖'),
], [2600, 1300, 5460])

doc.add_heading('14. 明确待办与下一阶段', level=1)
add_status('将底部绝对日期替换或组合为“多久以前 + 具体地点”。', '规划中')
add_status('根据下一张照片宽高比提前预加载并驱动卡片容器连续形变。', '规划中')
add_status('实现第二段“洗牌拼贴模式”：横屏、竖屏、方形混合的不规则布局。', '规划中')
add_status('拼贴每次重算大小与位置，并以日期为视觉锚点。', '规划中')
add_status('拼贴模式增加隐式位置刻度、双指调整内容大小与混入式教程卡。', '规划中')
add_status('建立后台记录 24 小时/7 天功耗、飞机、高铁及弱网地图加载基准。', '需真机验证')
add_status('用非重要照片验证系统删除、iCloud 同步及“最近删除”行为。', '需真机验证')

doc.add_heading('15. 核心验收清单', level=1)
for item in [
    '地图连续拖动、缩放时，路线、点和照片标记与地理位置保持一致，无屏幕层漂移。',
    '横屏与竖屏照片切换均不超出窗口宽度，不出现右侧溢出或整页偏移。',
    '左右滑动、上滑暂存、双指进入当日、Live Photo 播放之间不存在手势锁死。',
    '每组最多 20 张，少于 20 张不补齐；“再来一组”仍在当前地点重抽。',
    '多图删除复核布局在 1、2、4、9、12、20 张时均无裁切、重叠或按钮溢出。',
    '删除前出现 iOS 系统确认；取消后照片与 App 记录保持不变；成功后两者同步消失。',
    '后台记录不会因低精度漂移持续落点，高铁和飞机只保留表达行程所需的关键点。',
    '自定义瓦片源重启后仍可切换，失败时有明确回退，不阻塞地图主界面。',
]:
    add_bullet(item)

doc.add_heading('附录 A：当前测试记录', level=1)
add_bullet('自动回归：FootprintTests 71/71 通过。')
add_bullet('模拟器：iPhone 17 Pro / iOS 26.5，横屏与竖屏照片窗口约束、结束页布局已验证。')
add_bullet('真机：Birch / iPhone 15 Pro Max / iOS 26.0.1，已完成签名安装与启动。')
add_bullet('系统照片真实删除不做自动化破坏性测试，需由用户选择非重要照片手动确认。')

doc.add_heading('附录 B：关键技术边界', level=1)
add_bullet('Google 地图瓦片不得使用未经授权的非官方直链；自定义地图源必须满足服务条款和版权署名要求。')
add_bullet('iOS 后台定位由系统调度，无法保证固定秒级频率；产品应以路径表达、功耗和精度平衡为目标。')
add_bullet('系统照片删除确认无法由 App 跳过或完全仿制；这是用户数据安全边界。')
add_bullet('iCloud 删除会同步到登录同一 Apple 账户的设备，应在产品文案中持续明确提醒。')

doc.core_properties.title = '一生足迹｜产品设计说明书'
doc.core_properties.subject = '当前产品设计、实现状态与后续规划汇总'
doc.core_properties.author = 'Codex × 一生足迹产品团队'
doc.core_properties.keywords = '足迹, 地图, 照片, Live Photo, iOS, 产品文档'
doc.save(OUT)
print(OUT)
